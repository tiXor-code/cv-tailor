"""Three CV variants per package, all rendered from ONE set of tailored fields.

- technical     cv.pdf (+ cv.docx): ATS-safe single column; the default upload.
- designed      cv_designed.pdf: human-facing, a band of metric tiles on top.
                For email and referrals only, never for portals.
- designed_ats  cv_designed_ats.pdf (+ .docx): the designed content re-housed
                linearly; tile numbers become sentences, no band.

Metric tiles are chosen deterministically from the bullets that actually
render, so every tile number is verbatim profile content AND appears again in
a bullet below the band (nothing is lost if a parser strips the band).
validate_tiles() re-checks both properties before anything is drawn.

pick_variant() decides which file to send for a job and channel.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from cv_tailor.render import AcronymExpander, render_html, render_pdf
from cv_tailor.validate import validate_tiles

TECHNICAL_PDF = "cv.pdf"
TECHNICAL_DOCX = "cv.docx"
DESIGNED_PDF = "cv_designed.pdf"
DESIGNED_ATS_PDF = "cv_designed_ats.pdf"
DESIGNED_ATS_DOCX = "cv_designed_ats.docx"

MAX_TILES = 4
DESIGNED_MAX_PAGES = 2

# Senior titles get the designed-ATS twin on portals: same ATS safety, but it
# leads with the results band folded into sentences.
SENIOR_TITLE_RE = re.compile(r"\b(lead|head|principal|manager|architect)\b", re.IGNORECASE)

# A metric: optional sign/approx, optional currency, a number with optional
# thousands/decimals, then an optional %, K/M suffix, "/N" ratio or "of N".
_METRIC_RE = re.compile(
    r"(?<![\w/.\-+$#~])"
    r"(?P<value>[+~]?\$?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:%|[KM](?![A-Za-z]))?(?:/\d+(?:\.\d+)?)?(?:\s+of\s+\d+(?:,\d{3})*)?)"
    r"(?![\w\-/+.]*[\w/+])"  # reject 24GB, 21-agent, 2026-03-18, 94+/96+, 175K/mo
)
_STOP_TAIL = {"and", "with", "including", "at", "on", "in", "of", "the", "to",
              "for", "by", "a", "an", "from", "across", "per", "via", "while"}
_LABEL_MAX_WORDS = 4
# "200 OK" is a status code, not an achievement.
_JUNK_LABEL_START = {"ok"}


def _norm(text: str) -> str:
    return " ".join(str(text).split())


def rendered_bullets(profile: dict, fields: dict) -> list[tuple[str, str]]:
    """(source, bullet) for every bullet the CV renders, in render order."""
    exps = {e["id"]: e for e in profile.get("experiences", [])}
    projs = {p["id"]: p for p in profile.get("projects", [])}
    out: list[tuple[str, str]] = []
    for exp_id in fields.get("experience_ids_ordered", []):
        exp = exps.get(exp_id)
        if not exp:
            continue
        bullets = exp.get("bullets", [])
        for idx in fields.get("experience_bullets", {}).get(exp_id, []):
            if isinstance(idx, int) and 0 <= idx < len(bullets):
                out.append((exp.get("company", ""), _norm(bullets[idx])))
    for pid in fields.get("project_ids", []):
        proj = projs.get(pid)
        if not proj:
            continue
        for b in proj.get("bullets", []):
            out.append((proj.get("name", ""), _norm(b)))
    return out


def _label_after(text: str, end: int) -> str:
    """Up to a few words following the number, stopping at punctuation."""
    tail = text[end:]
    stop = re.search(r"[,.;:()\[\]]|\s[-–—]\s|->", tail)
    if stop:
        tail = tail[:stop.start()]
    words = tail.split()[:_LABEL_MAX_WORDS]
    while words and words[-1].lower() in _STOP_TAIL:
        words.pop()
    return " ".join(words)


def _label_before(text: str, start: int) -> str:
    """Up to a few words preceding the number, back to the last punctuation.
    Used when the words after it are only a preposition phrase ("Cut
    onboarding time by 35% across three teams" -> "Cut onboarding time")."""
    head = text[:start]
    cut = max(head.rfind(c) for c in ",.;:()[]")
    if cut >= 0:
        head = head[cut + 1:]
    words = head.split()[-_LABEL_MAX_WORDS:]
    while words and words[-1].lower() in _STOP_TAIL:
        words.pop()
    while words and words[0].lower() in _STOP_TAIL:
        words.pop(0)
    return " ".join(words)


def _label_for(text: str, m: re.Match) -> tuple[str, str]:
    """(label, phrase): the context words for the tile, and the verbatim span
    of the bullet covering number + label (used for the linear twin)."""
    after = _label_after(text, m.end())
    if after and after.split()[0].lower() not in _STOP_TAIL:
        end = text.find(after, m.end()) + len(after)
        return after, text[m.start():end]
    before = _label_before(text, m.start())
    if before:
        return before, text[text.rfind(before, 0, m.start()):m.end()]
    if after:
        end = text.find(after, m.end()) + len(after)
        return after, text[m.start():end]
    return "", ""


def _score(value: str) -> int:
    if any(c in value for c in "%$/") or " of " in value:
        return 3
    digits = re.sub(r"[^\d.]", "", value.split(" of ")[0].split("/")[0])
    try:
        num = float(digits)
    except ValueError:
        return 0
    if value.rstrip().endswith(("K", "M")) or num >= 100:
        return 2
    return 1 if num >= 10 else 0


def _is_year(value: str) -> bool:
    return bool(re.fullmatch(r"(19|20)\d\d", value))


def pick_tiles(profile: dict, fields: dict, max_tiles: int = MAX_TILES) -> list[dict]:
    """Deterministic metric tiles: the strongest number per rendered bullet,
    ranked by kind (ratio/percent/currency first, then magnitude), ties broken
    by render order. Each tile: value, label (context words next to the
    number in that same bullet), phrase (the verbatim bullet span holding
    both), source (company or project name)."""
    candidates = []
    order = 0
    for source, bullet in rendered_bullets(profile, fields):
        best = None
        for m in _METRIC_RE.finditer(bullet):
            value = m.group("value")
            if _is_year(value):
                continue
            label, phrase = _label_for(bullet, m)
            if not re.search(r"[A-Za-z]", label) or label.split()[0].lower() in _JUNK_LABEL_START:
                continue
            cand = (_score(value), order, {"value": value, "label": label,
                                       "phrase": phrase, "source": source})
            order += 1
            if best is None or cand[0] > best[0]:
                best = cand
        if best is not None:
            candidates.append(best)
    candidates.sort(key=lambda c: (-c[0], c[1]))
    # Two passes: first at most one tile per source (a band of four numbers
    # from one project reads as padding), then fill any free slots.
    chosen: list[tuple[int, dict]] = []
    seen, sources = set(), set()
    for one_per_source in (True, False):
        for _, order_idx, tile in candidates:
            key = (tile["value"], tile["label"].lower())
            if key in seen or (one_per_source and tile["source"] in sources):
                continue
            if len(chosen) >= max_tiles:
                break
            seen.add(key)
            sources.add(tile["source"])
            chosen.append((order_idx, tile))
    return [t for _, t in chosen]


def highlight_sentences(tiles: list[dict]) -> list[str]:
    """The tile band folded into plain sentences for the linear twin."""
    return [f"{t['source']}: {t['phrase']}." for t in tiles]


def pick_variant(entry: dict, channel: str, pkg_dir: Path | str | None = None) -> str:
    """File name to send for this job. channel "portal" -> the technical CV,
    or the designed-ATS twin for lead/head/principal/manager/architect titles;
    "email" -> the designed CV. With pkg_dir, a missing file falls back to
    cv.pdf so a caller never gets a name that does not exist."""
    if channel == "email":
        choice = DESIGNED_PDF
    elif channel == "portal":
        title = (entry or {}).get("title") or ""
        choice = DESIGNED_ATS_PDF if SENIOR_TITLE_RE.search(title) else TECHNICAL_PDF
    else:
        raise ValueError(f"unknown channel {channel!r} (expected 'portal' or 'email')")
    if pkg_dir is not None and not (Path(pkg_dir) / choice).exists():
        return TECHNICAL_PDF
    return choice


# ---------------------------------------------------------------- DOCX

def build_docx(profile: dict, fields: dict, out_path: Path | str,
               highlights: list[str] | None = None) -> Path:
    """Plain ATS-safe DOCX: real Heading styles, contact in the body (no
    header/footer), list bullets, no tables or text boxes."""
    from docx import Document
    from docx.shared import Cm, Pt

    expand = AcronymExpander()
    contact = profile.get("contact", {}) or {}
    name = contact.get("name", "")
    exps = {e["id"]: e for e in profile.get("experiences", [])}
    projs = {p["id"]: p for p in profile.get("projects", [])}

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    now = datetime.now(timezone.utc).replace(microsecond=0)
    cp = doc.core_properties
    cp.title = f"{name} - CV"
    cp.author = name
    cp.last_modified_by = name
    cp.subject = ""
    cp.keywords = ""
    cp.comments = ""
    cp.category = ""
    cp.revision = 1
    cp.created = now
    cp.modified = now

    doc.add_heading(name, level=0)
    parts = [contact.get(k) for k in ("location", "email", "phone", "website", "linkedin", "github")]
    doc.add_paragraph(" | ".join(str(p) for p in parts if p))

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(expand(_norm(fields.get("summary_rewrite", ""))))

    if highlights:
        doc.add_heading("Selected Results", level=1)
        for h in highlights:
            doc.add_paragraph(h, style="List Bullet")

    doc.add_heading("Experience", level=1)
    for exp_id in fields.get("experience_ids_ordered", []):
        exp = exps[exp_id]
        doc.add_heading(f"{exp.get('company', '')} - {exp.get('role', '')}", level=2)
        doc.add_paragraph(" | ".join(str(x) for x in (exp.get("dates"), exp.get("location")) if x))
        for idx in fields.get("experience_bullets", {}).get(exp_id, []):
            doc.add_paragraph(expand(_norm(exp["bullets"][idx])), style="List Bullet")

    if fields.get("project_ids"):
        doc.add_heading("Projects", level=1)
        for pid in fields["project_ids"]:
            proj = projs[pid]
            doc.add_heading(f"{proj.get('name', '')} - {', '.join(proj.get('tech', []))}", level=2)
            meta = expand(_norm(proj.get("tagline", "")))
            if proj.get("link"):
                meta = f"{meta} | {proj['link']}" if meta else proj["link"]
            if meta:
                doc.add_paragraph(meta)
            for b in proj.get("bullets", []):
                doc.add_paragraph(expand(_norm(b)), style="List Bullet")

    doc.add_heading("Skills", level=1)
    skills = profile.get("skills", {}) or {}
    emphasis = set(fields.get("skills_emphasis", []))
    groups = fields.get("skills_groups") or list(skills.keys())
    for group in groups:
        if group not in skills:
            continue
        para = doc.add_paragraph()
        para.add_run(f"{group.title()}: ").bold = True
        items = skills[group]
        for i, item in enumerate(items):
            run = para.add_run(item)
            run.bold = item in emphasis
            if i < len(items) - 1:
                para.add_run(", ")

    doc.add_heading("Education", level=1)
    for ed in profile.get("education", []) or []:
        doc.add_paragraph(f"{ed.get('degree', '')} - {ed.get('institution', '')} | {ed.get('year', '')}")
        if ed.get("notes"):
            doc.add_paragraph(_norm(ed["notes"]))

    out_path = Path(out_path)
    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------- orchestration

def pdf_page_count(path: Path | str) -> int | None:
    """Page count via pdfinfo; None when poppler is unavailable."""
    if not shutil.which("pdfinfo"):
        return None
    try:
        out = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True,
                             check=True, timeout=30).stdout
    except (subprocess.SubprocessError, OSError):
        return None
    m = re.search(r"^Pages:\s+(\d+)", out, re.MULTILINE)
    return int(m.group(1)) if m else None


def render_variants(profile: dict, fields: dict, entry: dict, pkg_dir: Path | str,
                    templates_dir: Path | str, *, pdf_renderer=render_pdf) -> dict:
    """Write cv.docx, cv_designed.pdf, cv_designed_ats.pdf/.docx next to the
    already-rendered cv.pdf. Each variant fails independently: an error is
    recorded under "errors" and pick_variant falls back to cv.pdf. Returns the
    meta.json "variants" block."""
    pkg_dir = Path(pkg_dir)
    templates_dir = Path(templates_dir)
    errors: list[str] = []
    warnings: list[str] = []
    produced: dict[str, dict] = {"technical": {"pdf": TECHNICAL_PDF}}

    tiles = pick_tiles(profile, fields)
    tile_errors = validate_tiles(profile, tiles, [b for _, b in rendered_bullets(profile, fields)])
    if tile_errors:
        errors.extend(f"tiles: {e}" for e in tile_errors)
        tiles = []  # never draw an unverified number
    highlights = highlight_sentences(tiles)

    def attempt(label, fn):
        try:
            fn()
            return True
        except Exception as exc:  # noqa: BLE001 -- one variant must not sink the package
            msg = f"{label}: {type(exc).__name__}: {exc}"
            errors.append(msg)
            print(f"cv variant failed: {msg}", file=sys.stderr)
            return False

    if attempt("technical.docx", lambda: build_docx(profile, fields, pkg_dir / TECHNICAL_DOCX)):
        produced["technical"]["docx"] = TECHNICAL_DOCX

    css = templates_dir / "cv_designed.css"

    def designed():
        out = pkg_dir / DESIGNED_PDF
        for compact in (False, True):
            html = render_html(profile, fields, templates_dir, "cv_designed.html.j2",
                               tiles=tiles, highlights=[], ats=False, compact=compact)
            pdf_renderer(html, css_path=css, out_path=out)
            pages = pdf_page_count(out)
            if pages is None or pages <= DESIGNED_MAX_PAGES:
                break
        if pages is not None and pages > DESIGNED_MAX_PAGES:
            warnings.append(f"{DESIGNED_PDF} runs to {pages} pages (max {DESIGNED_MAX_PAGES})")

    if attempt("designed.pdf", designed):
        produced["designed"] = {"pdf": DESIGNED_PDF}

    def designed_ats():
        html = render_html(profile, fields, templates_dir, "cv_designed.html.j2",
                           tiles=[], highlights=highlights, ats=True, compact=False)
        pdf_renderer(html, css_path=css, out_path=pkg_dir / DESIGNED_ATS_PDF)

    twin = {}
    if attempt("designed_ats.pdf", designed_ats):
        twin["pdf"] = DESIGNED_ATS_PDF
    if attempt("designed_ats.docx", lambda: build_docx(profile, fields, pkg_dir / DESIGNED_ATS_DOCX,
                                                       highlights=highlights)):
        twin["docx"] = DESIGNED_ATS_DOCX
    if twin:
        produced["designed_ats"] = twin

    block = dict(produced)
    block["tiles"] = tiles
    block["default"] = {
        "portal": pick_variant(entry, "portal", pkg_dir),
        "email": pick_variant(entry, "email", pkg_dir),
    }
    if warnings:
        block["warnings"] = warnings
    if errors:
        block["errors"] = errors
    return block
