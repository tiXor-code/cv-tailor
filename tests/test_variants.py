"""CV variants: technical (cv.pdf + cv.docx), designed (cv_designed.pdf) and
the designed-ATS twin (cv_designed_ats.pdf + .docx), all from one set of
tailored fields. Uses the fictional tests/fixtures/profile_variants.yaml and
renders real PDFs with WeasyPrint; pdftotext/pdfinfo assertions skip when
poppler is absent."""
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from cv_tailor import assemble as assemble_mod
from cv_tailor.assemble import assemble_package
from cv_tailor.profile import load_profile
from cv_tailor.render import AcronymExpander, render_html, render_pdf
from cv_tailor.validate import validate_tiles
from cv_tailor.variants import (
    build_docx, pdf_page_count, pick_tiles, pick_variant, render_variants, rendered_bullets,
)

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
FIXTURE = ROOT / "tests" / "fixtures" / "profile_variants.yaml"
HAS_POPPLER = bool(shutil.which("pdftotext") and shutil.which("pdfinfo"))
needs_poppler = pytest.mark.skipif(not HAS_POPPLER, reason="poppler (pdftotext/pdfinfo) not installed")

FIELDS = {
    "job_meta": {"company": "Initech", "role": "Platform Engineer", "location": "Remote",
                 "jd_url": None, "seniority_signal": "mid"},
    "chosen_summary_id": "default",
    "summary_rewrite": "Platform engineer who ships LLM automation and RAG lookups.",
    "experience_ids_ordered": ["acme", "globex"],
    "experience_bullets": {"acme": [0, 1, 2, 3], "globex": [0, 1]},
    "project_ids": ["widgetctl"],
    "skills_emphasis": ["Python"],
    "jd_keywords_matched": ["python"],
    "gaps_honest": [],
    "one_line_pitch": "A fit.",
}


@pytest.fixture
def profile():
    return load_profile(FIXTURE)


@pytest.fixture
def fields():
    return json.loads(json.dumps(FIELDS))


def _pdftotext(path):
    return subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True,
                          text=True, check=True).stdout


def _pdfinfo(path):
    out = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, check=True).stdout
    return {k.strip(): v.strip() for k, v in (l.split(":", 1) for l in out.splitlines() if ":" in l)}


def _in_order(text, needles):
    pos = [text.find(n) for n in needles]
    assert all(p >= 0 for p in pos), dict(zip(needles, pos))
    assert pos == sorted(pos), dict(zip(needles, pos))


@pytest.fixture
def package(tmp_path, profile, fields):
    """cv.pdf rendered as assemble does, then every other variant."""
    html = render_html(profile, fields, template_dir=TEMPLATES)
    render_pdf(html, css_path=TEMPLATES / "cv.css", out_path=tmp_path / "cv.pdf")
    block = render_variants(profile, fields, {"title": "Platform Engineer"}, tmp_path, TEMPLATES)
    return tmp_path, block


# ---------------------------------------------------------------- tiles

def test_pick_tiles_is_deterministic_and_verbatim(profile, fields):
    tiles = pick_tiles(profile, fields)
    assert tiles == pick_tiles(profile, fields)
    values = [t["value"] for t in tiles]
    # strongest per source first (percent / "of N"), then magnitude
    assert values == ["35%", "18 of 20", "212", "2,400"]
    by_value = {t["value"]: t for t in tiles}
    # "35% across three regional teams" is only a preposition phrase, so the
    # label comes from the words before the number
    assert by_value["35%"]["label"] == "Cut onboarding time"
    assert by_value["35%"]["phrase"] == "Cut onboarding time by 35%"
    assert by_value["212"]["label"] == "tests across 7 modules"
    assert by_value["18 of 20"]["source"] == "Globex Labs"


def test_pick_tiles_skips_years_hyphenated_units_and_unselected_bullets(profile, fields):
    values = {t["value"] for t in pick_tiles(profile, fields, max_tiles=20)}
    assert "2022" not in values          # a year is not a metric
    assert not {"4", "24"} & values      # 4-person, 24GB
    assert "777" not in values           # bullet not selected by the tailor


def test_every_tile_number_also_appears_in_a_rendered_bullet(profile, fields):
    tiles = pick_tiles(profile, fields)
    bullets = [b for _, b in rendered_bullets(profile, fields)]
    for t in tiles:
        assert any(t["value"] in b and t["phrase"] in b for b in bullets), t
    assert validate_tiles(profile, tiles, bullets) == []


def test_validate_tiles_rejects_numbers_not_in_profile(profile, fields):
    bullets = [b for _, b in rendered_bullets(profile, fields)]
    errors = validate_tiles(profile, [{"value": "99%", "label": "uptime", "source": "Acme Widgets"}], bullets)
    assert errors and "not found verbatim" in errors[0]


def test_validate_tiles_rejects_profile_number_missing_from_rendered_bullets(profile, fields):
    bullets = [b for _, b in rendered_bullets(profile, fields)]
    errors = validate_tiles(profile, [{"value": "777", "label": "widgets", "source": "Globex Labs"}], bullets)
    assert errors and "rendered bullet" in errors[0]


# ---------------------------------------------------------------- acronyms

def test_acronym_expander_expands_once_per_render():
    exp = AcronymExpander()
    assert exp("Built an LLM tool") == "Built an large language model (LLM) tool"
    assert exp("Another LLM and LLMs") == "Another LLM and LLMs"
    assert exp("KPIs met") == "key performance indicators (KPIs) met"
    assert exp("one KPI") == "one KPI"
    assert AcronymExpander()("ragged RAGs") == "ragged RAGs"  # whole tokens only


def test_technical_html_expands_acronym_once(profile, fields):
    html = render_html(profile, fields, template_dir=TEMPLATES)
    assert html.count("large language model (LLM)") == 1
    assert html.count("retrieval-augmented generation (RAG)") == 1


# ---------------------------------------------------------------- package

def test_render_variants_writes_every_file_and_records_defaults(package):
    pkg, block = package
    for name in ("cv.pdf", "cv.docx", "cv_designed.pdf", "cv_designed_ats.pdf", "cv_designed_ats.docx"):
        assert (pkg / name).exists(), name
    assert block["technical"] == {"pdf": "cv.pdf", "docx": "cv.docx"}
    assert block["designed"] == {"pdf": "cv_designed.pdf"}
    assert block["designed_ats"] == {"pdf": "cv_designed_ats.pdf", "docx": "cv_designed_ats.docx"}
    assert block["default"] == {"portal": "cv.pdf", "email": "cv_designed.pdf"}
    assert "errors" not in block
    assert [t["value"] for t in block["tiles"]] == ["35%", "18 of 20", "212", "2,400"]


@needs_poppler
def test_technical_pdf_is_ats_shaped_and_one_page(package):
    pkg, _ = package
    text = _pdftotext(pkg / "cv.pdf")
    head = "\n".join(text.splitlines()[:5])
    assert "alex@example.com" in head
    _in_order(text, ["Alex Example", "Summary", "Experience", "Acme Widgets", "Globex Labs",
                     "Projects", "Skills", "Education"])
    assert "Selected Results" not in text
    assert pdf_page_count(pkg / "cv.pdf") == 1


@needs_poppler
def test_designed_pdf_tile_band_precedes_body_and_numbers_repeat_below(package):
    pkg, block = package
    text = _pdftotext(pkg / "cv_designed.pdf")
    summary_at = text.find("Summary")
    experience_at = text.find("Experience")
    assert 0 < summary_at < experience_at
    for t in block["tiles"]:
        band_at = text.find(t["value"])
        assert 0 <= band_at < summary_at, f"tile {t['value']} not in the band"
        assert text.find(t["value"], experience_at) > experience_at, f"tile {t['value']} not repeated in a bullet"
    assert pdf_page_count(pkg / "cv_designed.pdf") <= 2


@needs_poppler
def test_designed_ats_twin_is_linear_with_results_as_sentences(package):
    pkg, block = package
    text = _pdftotext(pkg / "cv_designed_ats.pdf")
    head = "\n".join(text.splitlines()[:5])
    assert "alex@example.com" in head
    _in_order(text, ["Alex Example", "Summary", "Selected Results", "Experience", "Projects",
                     "Skills", "Education"])
    results = text[text.find("Selected Results"):text.find("Experience")]
    for t in block["tiles"]:
        assert t["value"] in results
    assert "Acme Widgets: Cut onboarding time by 35%." in " ".join(results.split())
    assert "large language model (LLM)" in text
    assert pdf_page_count(pkg / "cv_designed_ats.pdf") <= 2


@needs_poppler
@pytest.mark.parametrize("name", ["cv.pdf", "cv_designed_ats.pdf"])
def test_ats_variants_pass_the_ats_sanity_checks(package, profile, fields, name):
    from cv_tailor.ats_check import extract_text, run_checks

    pkg, _ = package
    warnings = run_checks(
        extract_text(pkg / name), profile, fields,
        experiences_by_id={e["id"]: e for e in profile["experiences"]},
        projects_by_id={p["id"]: p for p in profile["projects"]},
    )
    assert warnings == []


@needs_poppler
def test_pdf_metadata_comes_from_profile(package):
    pkg, _ = package
    for name in ("cv.pdf", "cv_designed.pdf", "cv_designed_ats.pdf"):
        info = _pdfinfo(pkg / name)
        assert info.get("Title") == "Alex Example - CV", name
        assert info.get("Author") == "Alex Example", name


def _docx_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml")}


@pytest.mark.parametrize("name", ["cv.docx", "cv_designed_ats.docx"])
def test_docx_is_plain_ats_safe(package, name):
    from docx import Document

    pkg, _ = package
    doc = Document(str(pkg / name))
    assert doc.tables == []
    parts = _docx_parts(pkg / name)
    body = parts["word/document.xml"]
    for marker in ("<w:tbl>", "txbx", "<w:drawing", "<v:shape", "w:framePr"):
        assert marker not in body, marker
    for sec in doc.sections:
        assert all(not p.text.strip() for p in sec.header.paragraphs)
        assert all(not p.text.strip() for p in sec.footer.paragraphs)

    paras = doc.paragraphs
    assert paras[0].text == "Alex Example"
    assert "alex@example.com" in paras[1].text and "+1 555 010 0199" in paras[1].text
    h1 = [p.text for p in paras if p.style.name == "Heading 1"]
    expected = ["Summary", "Experience", "Projects", "Skills", "Education"]
    if name == "cv_designed_ats.docx":
        expected.insert(1, "Selected Results")
    assert h1 == expected
    assert any(p.style.name == "List Bullet" for p in paras)

    cp = doc.core_properties
    assert cp.title == "Alex Example - CV"
    assert cp.author == "Alex Example"
    assert cp.last_modified_by == "Alex Example"
    assert "python-docx" not in (cp.comments or "") + (cp.subject or "") + (cp.keywords or "")


def test_docx_bullets_carry_every_tile_number(package):
    from docx import Document

    pkg, block = package
    bullets = [p.text for p in Document(str(pkg / "cv.docx")).paragraphs if p.style.name == "List Bullet"]
    for t in block["tiles"]:
        assert any(t["value"] in b for b in bullets), t


def test_variant_failure_is_recorded_and_default_falls_back(tmp_path, profile, fields):
    def broken_renderer(html, css_path, out_path):
        raise RuntimeError("renderer down")

    block = render_variants(profile, fields, {"title": "Head of Platform"}, tmp_path, TEMPLATES,
                            pdf_renderer=broken_renderer)
    assert "designed" not in block
    assert block["designed_ats"] == {"docx": "cv_designed_ats.docx"}
    assert any("designed.pdf" in e for e in block["errors"])
    assert block["default"] == {"portal": "cv.pdf", "email": "cv.pdf"}


# ---------------------------------------------------------------- pick_variant

@pytest.mark.parametrize("title,expected", [
    ("Platform Engineer", "cv.pdf"),
    ("Senior Backend Engineer", "cv.pdf"),
    ("Engineering Lead", "cv_designed_ats.pdf"),
    ("Head of Data", "cv_designed_ats.pdf"),
    ("Principal Engineer", "cv_designed_ats.pdf"),
    ("Engineering Manager", "cv_designed_ats.pdf"),
    ("Solutions Architect", "cv_designed_ats.pdf"),
    ("Leading Widgets Engineer", "cv.pdf"),  # whole words only
])
def test_pick_variant_portal(title, expected):
    assert pick_variant({"title": title}, "portal") == expected


def test_pick_variant_email_and_unknown_channel():
    assert pick_variant({"title": "Engineer"}, "email") == "cv_designed.pdf"
    with pytest.raises(ValueError):
        pick_variant({"title": "Engineer"}, "fax")


# ---------------------------------------------------------------- assemble wiring

def test_assemble_package_records_variants_in_meta(tmp_path, monkeypatch):
    monkeypatch.setenv("CV_TAILOR_PROFILE", str(FIXTURE))
    monkeypatch.setenv("CV_TAILOR_TEMPLATES", str(TEMPLATES))
    monkeypatch.setattr(assemble_mod, "tailor", lambda *a, **k: json.loads(json.dumps(FIELDS)))
    monkeypatch.setattr(assemble_mod, "cover_letter", lambda *a, **k: " ".join(["shipped"] * 140))
    monkeypatch.setattr(assemble_mod, "check_cover", lambda letter: [])

    def fake_pdf(html, css_path, out_path):
        Path(out_path).write_bytes(b"%PDF-1.4 fake\n")
        return Path(out_path)

    monkeypatch.setattr(assemble_mod, "render_pdf", fake_pdf)
    queue = tmp_path / "queue"
    (queue / "2026-01-01").mkdir(parents=True)
    (queue / "2026-01-01" / "descriptions.json").write_text(json.dumps({"j1": "Build platforms."}))
    entry = {"id": "j1", "title": "Staff Architect", "company": "Initech", "source": "lever",
             "url": "https://initech.example/jobs/1"}

    result = assemble_package(entry, "2026-01-01", queue_dir=queue, client=object())
    pkg = Path(result["package_dir"])
    meta = json.loads((pkg / "meta.json").read_text())
    v = meta["variants"]
    assert v["default"] == {"portal": "cv_designed_ats.pdf", "email": "cv_designed.pdf"}
    assert result["cv_path"].endswith("cv.pdf")
    for name in ("cv.pdf", "cv.docx", "cv_designed.pdf", "cv_designed_ats.pdf", "cv_designed_ats.docx"):
        assert (pkg / name).exists(), name
